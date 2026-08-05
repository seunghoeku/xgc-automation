"""Promote archived inputs toward validated kb/cases/ baselines.

The input archive holds hundreds of historical inputs, but an archived file
alone does not say whether the simulation worked. The researcher's simulation
notebook (``Other_data/simulation_note_XGC1.xlsx``) records the intent and the
outcome of each numbered run, so promotion decisions join both sources:

    input_backup/<file>   -> what was configured (parameters, schema age)
    simulation notebook   -> why it was run and how it ended

Everything here is deterministic and evidence-quoting: outcomes are classified
by keyword rules over the researcher's own words, never invented, and the
original note text is always carried along so a human can check the reading.

A promoted case starts as a *draft* under ``kb/cases/drafts/``. Drafts are
invisible to ``input_builder.list_cases()`` (and therefore to the physics
advisor), because an input only becomes a baseline after its referenced files
are staged and the current ``xgc --dryrun`` accepts it (docs/input-archive.md).
Finalizing a draft into ``kb/cases/`` stays a human step.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml

from tools import input_archive

PROJECT_ROOT = Path(__file__).resolve().parent.parent
NOTES_PATH = PROJECT_ROOT / "Other_data" / "simulation_note_XGC1.xlsx"
DRAFTS_DIR = PROJECT_ROOT / "kb" / "cases" / "drafts"

_SHOT_NUMBER_RE = re.compile(r"(?:^|[^0-9])(\d{2,4})(?:[^0-9]|$)")

# Header names vary across the notebook's yearly sheets; normalize them.
_HEADER_ALIASES = {
    "shot #": "shot",
    "shotname": "shot",
    "short description": "description",
    "chracteristic": "description",   # sic, as written in the notebook
    "characteristic": "description",
    "code": "code",
    "code version": "code",
    "geometry": "geometry",
    "note": "note",
    "general description": "note",
    "simulation date": "date",
    "data": "data",
}

# Deterministic outcome keywords, matched case-insensitively against the
# researcher's own note/description text. Order matters only for reporting;
# any failure signal marks the run as failed even if success words appear.
_FAILURE_PATTERNS = [
    r"\bfail(?:ed|ure)?\b",
    r"\bwrong\b",
    r"\bnot\s+run\b",
    r"\bnumerical\s+instabilit",
    r"\binstabilit(?:y|ies)\s+(?:at|in|near|inboard)",
    r"\bout\s+of\s+memor",
    r"\bnan\b",
    r"\bproblem(?:atic)?\b",
    r"\bstop(?:ped)?\s+due\s+to\b",
]
_SUCCESS_PATTERNS = [
    r"\bok\b",
    r"\bused\s+for\s+paper\b",
    r"\bgood\b",
    r"\bworks\s+fine\b",
]
# Negations that turn a success word into its opposite ("not that good",
# "not good", "no good"). Checked before success words get credit.
_NEGATED_SUCCESS_PATTERNS = [
    r"\bnot\s+(?:that\s+|so\s+|very\s+)?good\b",
    r"\bno\s+good\b",
    r"\bnot\s+ok\b",
    r"\bdidn'?t\s+work\b",
    r"\bdoes\s*n[o']t\s+work\b",
]
_CAUTION_PATTERNS = [
    r"\bpreliminar",
    r"\btoo\s+(?:high|low|large|small)\b",
    r"\bnot\s+much\s+meaning",
    r"\bno\s+interesting\b",
    r"\btest\b",
]


class CasePromotionError(RuntimeError):
    """The promotion request cannot be satisfied safely."""


@dataclass
class SimulationNote:
    """One row of the researcher's simulation notebook."""

    shot: str
    number: int | None
    sheet: str
    description: str = ""
    code: str = ""
    geometry: str = ""
    note: str = ""
    date: str = ""
    data: str = ""

    def combined_text(self) -> str:
        return " ".join(part for part in (self.description, self.note) if part)

    def to_dict(self) -> dict[str, Any]:
        return {
            "shot": self.shot,
            "number": self.number,
            "sheet": self.sheet,
            "description": self.description,
            "code": self.code,
            "geometry": self.geometry,
            "note": self.note,
            "date": self.date,
            "data": self.data,
        }


def _shot_number(shot: str) -> int | None:
    match = _SHOT_NUMBER_RE.search(shot)
    return int(match.group(1)) if match else None


def load_simulation_notes(path: Path | None = None) -> list[SimulationNote]:
    """Parse the simulation notebook into normalized per-run records.

    The notebook is read-only evidence; this never writes to it. Sheets have
    drifted in layout over the years, so columns are mapped by header name
    rather than position.
    """
    notes_path = path or NOTES_PATH
    if not notes_path.is_file():
        raise CasePromotionError(f"simulation notebook not found: {notes_path}")

    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise CasePromotionError(
            "openpyxl is required to read the simulation notebook"
        ) from exc

    workbook = openpyxl.load_workbook(notes_path, read_only=True, data_only=True)
    notes: list[SimulationNote] = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            continue
        columns: dict[int, str] = {}
        for idx, cell in enumerate(header_row):
            if cell is None:
                continue
            key = _HEADER_ALIASES.get(str(cell).strip().lower())
            if key:
                columns[idx] = key

        if "shot" not in columns.values():
            continue

        for row in rows:
            record: dict[str, str] = {}
            for idx, key in columns.items():
                value = row[idx] if idx < len(row) else None
                if value is None:
                    continue
                text = str(value).strip()
                if text:
                    record[key] = text
            shot = record.get("shot", "")
            if not shot or shot.lower() == "empty":
                continue
            notes.append(
                SimulationNote(
                    shot=shot,
                    number=_shot_number(shot),
                    sheet=sheet_name,
                    description=record.get("description", ""),
                    code=record.get("code", ""),
                    geometry=record.get("geometry", ""),
                    note=record.get("note", ""),
                    date=record.get("date", ""),
                    data=record.get("data", ""),
                )
            )
    workbook.close()
    return notes


def notes_for_number(
    number: int, notes: list[SimulationNote] | None = None
) -> list[SimulationNote]:
    """All notebook rows whose shot number matches."""
    notes = notes if notes is not None else load_simulation_notes()
    return [note for note in notes if note.number == number]


def classify_outcome(text: str) -> dict[str, Any]:
    """Classify a run outcome from the researcher's own words.

    Returns the classification plus the exact phrases that produced it, so the
    reading is checkable. Failure evidence dominates: a note like "OK at first,
    failed later" is a failure.
    """
    def _hits(patterns: list[str]) -> list[str]:
        found: list[str] = []
        for pattern in patterns:
            match = re.search(pattern, text, flags=re.IGNORECASE)
            if match:
                found.append(match.group(0))
        return found

    negated = _hits(_NEGATED_SUCCESS_PATTERNS)
    failure = _hits(_FAILURE_PATTERNS)
    success = _hits(_SUCCESS_PATTERNS)
    caution = _hits(_CAUTION_PATTERNS)

    if negated:
        # "not that good" contains "good"; strip the false success credit and
        # treat the negation itself as caution-grade evidence, not failure --
        # the run completed, the physics just disappointed.
        success = []
        caution = negated + caution

    if failure:
        outcome = "failed"
    elif success:
        outcome = "success"
    elif caution:
        outcome = "inconclusive"
    else:
        outcome = "unknown"

    return {
        "outcome": outcome,
        "failure_evidence": failure,
        "success_evidence": success,
        "caution_evidence": caution,
        "note_text": text,
    }


# Device/experiment tokens for detecting a filename that disagrees with the
# notebook entry of the same number (e.g. an ASDEX-named file matched to a
# d3d note means one of the two records the wrong number).
_DEVICE_TOKENS = {
    "iter": {"iter", "jintrac"},
    "st40": {"st40"},
    "kstar": {"kstar", "kstart"},
    "d3d": {"d3d", "diii-d", "diiid"},
    "asdex": {"asdex", "aug", "au"},
    "tcv": {"tcv"},
    "nstx": {"nstx"},
    "cmod": {"cmod", "c-mod"},
    "jet": {"jet"},
    "west": {"west"},
    "cbc": {"cbc", "ecbc"},
}


def _device_of(text: str) -> str | None:
    words = {word.lower() for word in re.findall(r"[A-Za-z0-9-]+", text)}
    for device, tokens in _DEVICE_TOKENS.items():
        if words & tokens:
            return device
    return None


def name_note_mismatch(
    archive_name: str, matched: list[SimulationNote]
) -> dict[str, Any] | None:
    """Flag a device disagreement between the filename and its notebook entry.

    Returns None when either side names no known device, or when they agree.
    A mismatch does not say which side is wrong -- only that the join on the
    simulation number should not be trusted without a human look.
    """
    file_device = _device_of(archive_name)
    if file_device is None:
        return None
    for note in matched:
        note_device = _device_of(
            " ".join((note.description, note.geometry, note.note))
        )
        if note_device is not None and note_device != file_device:
            return {
                "filename_device": file_device,
                "notebook_device": note_device,
                "notebook_shot": note.shot,
                "notebook_description": note.description,
                "detail": (
                    f"filename says {file_device!r} but notebook entry "
                    f"{note.shot!r} says {note_device!r}; the number join "
                    f"may be wrong on either side"
                ),
            }
    return None


def _outcome_for_notes(matched: list[SimulationNote]) -> dict[str, Any]:
    if not matched:
        return {
            "outcome": "unknown",
            "failure_evidence": [],
            "success_evidence": [],
            "caution_evidence": [],
            "note_text": "",
        }
    return classify_outcome(" ".join(note.combined_text() for note in matched))


# The notebook's Code column records which binary era a run belongs to.
# Until recently EM and ES were separate executables (xgc-eem / xgc-es) and
# neoclassical runs used (and still use) xgca; the field model was chosen by
# the binary, not the input. Current XGC unifies EM/ES into one binary with
# the sml_electromagnetic switch (default .false. = electrostatic), so an
# old EM input MUST gain sml_electromagnetic=.true. during migration or it
# will silently run electrostatic.
_CODE_XGCA_RE = re.compile(r"\bxgca\b", re.IGNORECASE)
_CODE_EM_RE = re.compile(r"\bem\b", re.IGNORECASE)
_CODE_ES_RE = re.compile(r"\bes\b", re.IGNORECASE)


def classify_code_column(code_text: str) -> dict[str, Any]:
    """Interpret the notebook Code column into binary era facts.

    Returns flavor (xgc1/xgca/unknown) and field model (em/es/unknown),
    quoting the original text. XGCa runs are neoclassical: the field model
    distinction does not select their binary.
    """
    text = code_text.strip()
    if _CODE_XGCA_RE.search(text):
        flavor = "xgca"
    elif text:
        flavor = "xgc1"
    else:
        flavor = "unknown"

    if _CODE_EM_RE.search(text):
        field = "em"
    elif _CODE_ES_RE.search(text):
        field = "es"
    else:
        field = "unknown"

    return {
        "flavor": flavor,
        "field_model": field,
        "binary_era": (
            "xgca" if flavor == "xgca"
            else {"em": "xgc-eem", "es": "xgc-es"}.get(field, "unknown")
        ),
        "code_text": text,
    }


def _infer_physics(
    key_parameters: dict[str, str], matched_notes: list[SimulationNote]
) -> dict[str, Any]:
    """Deterministic physics classification with per-field sources.

    Every field records where it came from (a parameter value or the notebook
    Code column). Anything not derivable is reported as unknown rather than
    guessed.
    """
    physics: dict[str, Any] = {}

    # Field model. Historically the field model was chosen by the binary
    # (xgc-eem vs xgc-es), not the input, so an explicit input switch is rare
    # in archived files: sml_electromagnetic (current) or legacy sml_em.
    # Otherwise the notebook Code column is the only record of which binary
    # the run used.
    em_param = key_parameters.get("sml_electromagnetic") or key_parameters.get("sml_em")
    if em_param is not None:
        is_em = em_param.strip().lower() in {".true.", ".t.", "t", "true"}
        physics["field_model"] = {
            "value": "electromagnetic" if is_em else "electrostatic",
            "source": "input parameter (sml_electromagnetic/sml_em)",
        }
    else:
        for note in matched_notes:
            code = classify_code_column(note.code)
            if code["field_model"] != "unknown":
                physics["field_model"] = {
                    "value": (
                        "electromagnetic"
                        if code["field_model"] == "em"
                        else "electrostatic"
                    ),
                    "source": (
                        f"notebook Code column: {note.code!r} "
                        f"(binary era: {code['binary_era']}); "
                        f"not declared in the input itself"
                    ),
                }
                break
        else:
            physics["field_model"] = {
                "value": "unknown",
                "source": None,
            }

    # Flavor: xgca (neoclassical) still has its own binary; xgc1 is implied
    # otherwise. The Code column is the only record for archived runs.
    flavor = "unknown"
    flavor_source = None
    for note in matched_notes:
        code = classify_code_column(note.code)
        if code["flavor"] != "unknown":
            flavor = code["flavor"]
            flavor_source = f"notebook Code column: {note.code!r}"
            break
    physics["flavor"] = {"value": flavor, "source": flavor_source}

    # Domain: outpsi > 1 includes the separatrix/SOL; bounce settings agree.
    outpsi = key_parameters.get("sml_outpsi")
    if outpsi is not None:
        try:
            outside = float(outpsi.lower().replace("d", "e")) > 1.0
        except ValueError:
            outside = None
        if outside is None:
            physics["domain"] = {"value": "unknown", "source": None}
        else:
            physics["domain"] = {
                "value": "edge-including" if outside else "core",
                "source": f"sml_outpsi={outpsi}",
            }
    else:
        physics["domain"] = {"value": "unknown", "source": None}

    marker = key_parameters.get("ptl_marker_type")
    deltaf = key_parameters.get("sml_deltaf")
    if marker is not None:
        physics["marker_model"] = {"value": marker, "source": "ptl_marker_type"}
    elif deltaf is not None:
        physics["marker_model"] = {
            "value": f"legacy sml_deltaf={deltaf} (inert in current XGC)",
            "source": "sml_deltaf",
        }
    else:
        physics["marker_model"] = {
            "value": "TotalF (XGC default; nothing declared)",
            "source": "absence of ptl_marker_type",
        }

    col_mode = key_parameters.get("col_mode")
    physics["collisions"] = {
        "value": (
            "off" if col_mode == "0" else f"col_mode={col_mode}"
        ) if col_mode is not None else "unknown",
        "source": "col_mode" if col_mode is not None else None,
    }

    return physics


def _archive_record(archive_name: str) -> dict[str, Any]:
    inventory = input_archive.inspect_archive(min_number=0, max_files=2_000)
    for record in inventory["records"]:
        if record["archive_name"] == archive_name:
            return record
    raise CasePromotionError(
        f"{archive_name!r} not found in {inventory['archive_root']} "
        f"(is it a numbered input?)"
    )


def rank_promotion_candidates(
    *,
    min_number: int = 600,
    max_files: int = 100,
) -> dict[str, Any]:
    """Join the archive inventory with notebook outcomes and rank candidates.

    Ranking is deterministic: recorded successes first, then schema alignment,
    then recency. Runs the researcher recorded as failed are listed but never
    ranked above unknowns -- a failed configuration is evidence of what not to
    promote.
    """
    inventory = input_archive.inspect_archive(
        min_number=min_number, max_files=max_files
    )
    notes = load_simulation_notes()

    candidates: list[dict[str, Any]] = []
    for record in inventory["records"]:
        matched = notes_for_number(record["simulation_number"], notes)
        outcome = _outcome_for_notes(matched)
        mismatch = name_note_mismatch(record["archive_name"], matched)
        code = classify_code_column(matched[0].code) if matched else None
        candidates.append(
            {
                "archive_name": record["archive_name"],
                "simulation_number": record["simulation_number"],
                "variant": record.get("variant"),
                "name_note_mismatch": mismatch,
                "binary_era": code,
                "machines": record["machines"],
                "schema_alignment": record["schema_comparison"]["alignment"],
                "recognized_ratio": record["schema_comparison"]["recognized_ratio"],
                "outcome": outcome["outcome"],
                "outcome_evidence": {
                    key: outcome[key]
                    for key in (
                        "failure_evidence",
                        "success_evidence",
                        "caution_evidence",
                    )
                },
                "notebook_entries": [note.to_dict() for note in matched],
                "filename_tags": record["filename_tags"],
                "key_parameters": record["key_parameters"],
            }
        )

    outcome_rank = {"success": 0, "unknown": 1, "inconclusive": 2, "failed": 3}
    alignment_rank = {"high": 0, "partial": 1, "low": 2, "not_compared": 3}
    candidates.sort(
        key=lambda item: (
            outcome_rank.get(item["outcome"], 2),
            # Variants (.tar/.inc/.neo/backups) document a primary input but
            # are not independent runs; rank them below primaries.
            item["variant"] is not None,
            alignment_rank.get(item["schema_alignment"], 3),
            -item["simulation_number"],
        )
    )

    return {
        "schema": "xgc-case-promotion-candidates/v1",
        "notebook": str(NOTES_PATH),
        "selection": inventory["selection"],
        "candidates": candidates,
        "warnings": [
            "Outcomes come from the researcher's notebook via keyword rules; "
            "verify the quoted evidence before trusting a classification.",
            "Promotion still requires staging referenced files and passing "
            "xgc --dryrun; ranking alone validates nothing.",
        ],
    }


def plan_case_promotion(archive_name: str) -> dict[str, Any]:
    """Draft a kb case for one archived input. Writes nothing.

    The draft combines the parsed input, the notebook's intent/outcome, and a
    deterministic physics classification, and lists exactly what remains before
    the case can become a baseline.
    """
    record = _archive_record(archive_name)
    matched = notes_for_number(record["simulation_number"], load_simulation_notes())
    outcome = _outcome_for_notes(matched)
    physics = _infer_physics(record["key_parameters"], matched)
    mismatch = name_note_mismatch(record["archive_name"], matched)

    warnings: list[str] = []
    if record.get("variant"):
        warnings.append(
            f"This file is a {record['variant']!r} variant, not a primary "
            f"input; prefer promoting the primary and recording the variant "
            f"as context."
        )
    if mismatch:
        warnings.append("Name/notebook device mismatch: " + mismatch["detail"])

    # Binary-era trap: EM inputs from the xgc-eem era carry no field-model
    # switch, and the unified binary defaults to electrostatic. Without adding
    # sml_electromagnetic=.true. such a case would dryrun cleanly and then
    # silently simulate the wrong physics.
    field = physics.get("field_model") or {}
    if (
        field.get("value") == "electromagnetic"
        and "notebook" in (field.get("source") or "")
    ):
        warnings.append(
            "The notebook records this as an EM run from the xgc-eem binary "
            "era, but the input has no field-model switch. Migration MUST add "
            "sml_electromagnetic=.true. -- the unified binary defaults to "
            "electrostatic and would silently run the wrong physics."
        )
    if (physics.get("flavor") or {}).get("value") == "xgca":
        warnings.append(
            "This is an XGCa (neoclassical) run; it uses the separate xgca "
            "binary, which this pipeline has not validated yet."
        )
    if outcome["outcome"] == "failed":
        warnings.append(
            "The notebook records this run as FAILED: "
            + "; ".join(outcome["failure_evidence"])
            + ". Promoting a failed configuration is almost never what you want."
        )
    if outcome["outcome"] == "unknown":
        warnings.append(
            "No notebook entry (or no outcome wording) was found for this run; "
            "intent and result are unconfirmed."
        )
    if record["schema_comparison"]["alignment"] not in {"high", "not_compared"}:
        unknown = record["schema_comparison"]["unknown_parameters"]
        warnings.append(
            f"{len(unknown)} parameter(s) are not in current XGC metadata and "
            f"will need migration: {', '.join(unknown[:10])}"
            + ("..." if len(unknown) > 10 else "")
        )

    draft_name = f"draft-{archive_name}".lower()
    draft = {
        "name": draft_name,
        "status": "draft",
        "description": (
            matched[0].description if matched else f"promoted from {archive_name}"
        ),
        "promotion": {
            "archive_file": record["relative_path"],
            "archive_sha256": record["sha256"],
            "simulation_number": record["simulation_number"],
            "variant": record.get("variant"),
            "name_note_mismatch": mismatch,
            "machines": record["machines"],
            "notebook_entries": [note.to_dict() for note in matched],
            "outcome": outcome,
        },
        "physics_inferred": physics,
        "key_parameters": record["key_parameters"],
        "schema_comparison": record["schema_comparison"],
        "remaining_steps": [
            "Locate and stage the referenced mesh/equilibrium/profile files "
            "on the target machine.",
            "Run validate_input (xgc --dryrun) and migrate any removed or "
            "deprecated parameters.",
            "Record validated:/resource_estimate: from the dryrun, then move "
            "the draft from kb/cases/drafts/ into kb/cases/ with advisor "
            "match signals.",
        ],
    }

    return {
        "schema": "xgc-case-promotion-plan/v1",
        "archive_name": archive_name,
        "draft_name": draft_name,
        "draft_path": str(DRAFTS_DIR / f"{draft_name}.yaml"),
        "draft": draft,
        "outcome": outcome["outcome"],
        "warnings": warnings,
        "next_action": (
            "Review the draft and warnings; if acceptable, call "
            "write_draft_case with confirm=True."
        ),
    }


def write_draft_case(
    archive_name: str,
    *,
    confirm: bool = False,
    allow_failed: bool = False,
) -> dict[str, Any]:
    """Write a promotion draft under kb/cases/drafts/. Requires confirmation.

    Drafts never overwrite silently, and a run the notebook records as failed
    is refused unless ``allow_failed=True`` (a failed configuration can still
    be useful as documentation of what went wrong, but that must be a human
    decision).
    """
    plan = plan_case_promotion(archive_name)

    if plan["outcome"] == "failed" and not allow_failed:
        raise CasePromotionError(
            f"the notebook records {archive_name!r} as failed; "
            f"pass allow_failed=True to draft it anyway"
        )

    if not confirm:
        return {
            "written": False,
            "reason": "confirmation required; review the plan and resubmit "
            "with confirm=True",
            "plan": plan,
        }

    draft_path = Path(plan["draft_path"])
    if draft_path.exists():
        raise CasePromotionError(
            f"{draft_path} already exists; remove it first if you mean to redo "
            f"the draft"
        )

    DRAFTS_DIR.mkdir(parents=True, exist_ok=True)
    draft_path.write_text(
        yaml.safe_dump(plan["draft"], sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )

    return {
        "written": True,
        "draft_path": str(draft_path),
        "outcome": plan["outcome"],
        "warnings": plan["warnings"],
        "remaining_steps": plan["draft"]["remaining_steps"],
    }
